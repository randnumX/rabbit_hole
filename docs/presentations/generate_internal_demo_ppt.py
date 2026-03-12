from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT_PATH = Path(__file__).with_name("rabbithole-internal-demo.pptx")

BG = RGBColor(17, 24, 39)
PANEL = RGBColor(30, 41, 59)
TEXT = RGBColor(241, 245, 249)
MUTED = RGBColor(148, 163, 184)
ACCENT = RGBColor(245, 158, 11)
ACCENT_ALT = RGBColor(56, 189, 248)
GOOD = RGBColor(74, 222, 128)
WARN = RGBColor(251, 191, 36)
BAD = RGBColor(248, 113, 113)


SLIDES = [
    {
        "title": "RabbitHole",
        "subtitle": "Visualizing when a chat stays on track, branches out, or drifts",
        "bullets": [
            "Internal demo",
            "Chrome extension + FastAPI backend",
            "Still in development, but the interaction model is working",
        ],
        "kind": "title",
    },
    {
        "title": "Problem Statement",
        "bullets": [
            "Start with one core task",
            "Chat branches into tangents",
            "Some are useful, not urgent",
            "UI shows transcript, not structure",
        ],
    },
    {
        "title": "What I Wanted the Tool to Do",
        "bullets": [
            "Surface the main thread",
            "Separate tangent vs drift",
            "Return me to the core task",
            "Save rabbit holes for later",
        ],
    },
    {
        "title": "The General Direction",
        "bullets": [
            "Visual map of chat drift",
            "Main path, branches, breaks",
            "Interactive, not just analytical",
            "Built inside ChatGPT",
        ],
    },
    {
        "title": "Why a Rabbit and a Trail",
        "bullets": [
            "Trail = progress",
            "Branch = tangent",
            "Broken node = drift",
            "Return arc = recovery",
        ],
    },
    {
        "title": "How This Can Be Achieved",
        "bullets": [
            "Extract exchanges",
            "Normalize and summarize",
            "Score semantic continuity",
            "Classify each exchange",
            "Render the trail",
        ],
    },
    {
        "title": "System Architecture",
        "bullets": [
            "Content script",
            "Background worker",
            "FastAPI backend",
            "Shared types",
            "Sidebar UI",
        ],
        "kind": "architecture",
    },
    {
        "title": "How the Backend Classifies Drift",
        "bullets": [
            "Root topic from the opening turns",
            "Mainline plus branch lineages",
            "Similarity, continuity, drift, novelty",
            "Rule-ordered classification",
            "Output states drive the trail",
        ],
    },
    {
        "title": "Analysis Modes",
        "bullets": [
            "Deterministic",
            "Hybrid",
            "Probabilistic",
            "Same UI, different analysis depth",
        ],
        "kind": "modes",
    },
    {
        "title": "Live Demo",
        "bullets": [
            "Compact rabbit launcher",
            "Docked sidebar",
            "Exchange-level nodes",
            "Click node to jump to transcript",
            "Live rabbit states",
        ],
    },
    {
        "title": "Current Prototype State",
        "bullets": [
            "End-to-end flow works",
            "Extraction, scoring, rendering are live",
            "Click-to-scroll works",
            "Auto-analysis works",
            "Still refining edge cases",
        ],
    },
    {
        "title": "What's in the Pipeline",
        "bullets": [
            "Better topic labels",
            "Better branch summaries",
            "Better unrelated-topic detection",
            "Support beyond ChatGPT",
            "Save rabbit holes for later",
            "Better replay and navigation",
        ],
    },
    {
        "title": "End",
        "bullets": [
            "Make drift visible",
            "Stay on the main task",
            "Revisit side paths later",
        ],
        "kind": "end",
    },
]


def set_background(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_title(slide, title):
    box = slide.shapes.add_textbox(Inches(0.85), Inches(0.55), Inches(11.4), Inches(0.8))
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = "Avenir Next"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = TEXT


def add_footer(slide, idx, total):
    box = slide.shapes.add_textbox(Inches(0.85), Inches(7.0), Inches(11.4), Inches(0.3))
    frame = box.text_frame
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = f"{idx}/{total}"
    run.font.name = "Avenir Next"
    run.font.size = Pt(10)
    run.font.color.rgb = MUTED


def add_bullets(slide, bullets):
    card = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(1.45), Inches(7.2), Inches(5.15)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = PANEL
    card.line.color.rgb = RGBColor(51, 65, 85)
    card.adjustments[0] = 0.08

    tf = card.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(18)
    tf.margin_right = Pt(18)
    tf.margin_top = Pt(14)
    tf.vertical_anchor = MSO_ANCHOR.TOP

    for index, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.space_after = Pt(10)
        p.bullet = True
        p.font.name = "Avenir Next"
        p.font.size = Pt(20)
        p.font.color.rgb = TEXT


def add_side_note(slide, title, lines):
    note = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(8.35), Inches(1.45), Inches(3.45), Inches(5.15)
    )
    note.fill.solid()
    note.fill.fore_color.rgb = RGBColor(15, 23, 42)
    note.line.color.rgb = RGBColor(51, 65, 85)
    note.adjustments[0] = 0.08

    tf = note.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(16)
    tf.margin_top = Pt(16)

    head = tf.paragraphs[0]
    head.text = title
    head.font.name = "Avenir Next"
    head.font.size = Pt(15)
    head.font.bold = True
    head.font.color.rgb = ACCENT
    head.space_after = Pt(10)

    for line in lines:
        p = tf.add_paragraph()
        p.text = line
        p.font.name = "Avenir Next"
        p.font.size = Pt(13)
        p.font.color.rgb = MUTED
        p.space_after = Pt(8)


def add_architecture_visual(slide):
    labels = [
        ("ChatGPT DOM", ACCENT),
        ("Content Script", ACCENT_ALT),
        ("Background Worker", GOOD),
        ("FastAPI Backend", WARN),
        ("Trail UI", BAD),
    ]
    x = Inches(8.55)
    y = Inches(1.7)
    for idx, (label, color) in enumerate(labels):
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y + Inches(idx * 0.86), Inches(2.95), Inches(0.55)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.color.rgb = color
        tf = shape.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.name = "Avenir Next"
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42)


def add_modes_visual(slide):
    cards = [
        ("Deterministic", "Stable, tunable, explainable", ACCENT),
        ("Hybrid", "LLM helps summarize noisy replies", ACCENT_ALT),
        ("Probabilistic", "LLM overlay can move classification", GOOD),
    ]
    start_x = Inches(8.2)
    y = Inches(1.7)
    for idx, (title, subtitle, color) in enumerate(cards):
        card = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            start_x,
            y + Inches(idx * 1.45),
            Inches(3.1),
            Inches(0.95),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = PANEL
        card.line.color.rgb = color
        tf = card.text_frame
        tf.clear()
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Avenir Next"
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = color

        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = "Avenir Next"
        p2.font.size = Pt(12)
        p2.font.color.rgb = MUTED


def add_title_slide(slide, slide_data):
    box = slide.shapes.add_textbox(Inches(0.95), Inches(1.0), Inches(10.8), Inches(1.0))
    frame = box.text_frame
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = slide_data["title"]
    run.font.name = "Avenir Next"
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = TEXT

    sub = slide.shapes.add_textbox(Inches(0.95), Inches(1.95), Inches(10.6), Inches(0.7))
    tf = sub.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = slide_data["subtitle"]
    run.font.name = "Avenir Next"
    run.font.size = Pt(18)
    run.font.color.rgb = MUTED

    pill = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(3.0), Inches(3.0), Inches(0.5)
    )
    pill.fill.solid()
    pill.fill.fore_color.rgb = ACCENT
    pill.line.color.rgb = ACCENT
    tf = pill.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Internal technical demo"
    r.font.name = "Avenir Next"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(17, 24, 39)

    add_bullets(slide, slide_data["bullets"])


def add_end_slide(slide, bullets):
    add_title(slide, "End")
    add_bullets(slide, bullets)
    add_side_note(
        slide,
        "Close",
        [
            "Make drift visible.",
            "Stay on the main task.",
            "Come back to rabbit holes later.",
        ],
    )


def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]
    total = len(SLIDES)

    for idx, slide_data in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(blank)
        set_background(slide)

        kind = slide_data.get("kind")
        if kind == "title":
            add_title_slide(slide, slide_data)
        elif kind == "end":
            add_end_slide(slide, slide_data["bullets"])
        else:
            add_title(slide, slide_data["title"])
            add_bullets(slide, slide_data["bullets"])

            if kind == "architecture":
                add_architecture_visual(slide)
            elif kind == "modes":
                add_modes_visual(slide)
            else:
                add_side_note(
                    slide,
                    "Presenter cue",
                    [
                        "Keep this slide at a high level.",
                        "Use the live demo to make it concrete.",
                    ],
                )

        add_footer(slide, idx, total)

    prs.save(OUT_PATH)


if __name__ == "__main__":
    build_presentation()
